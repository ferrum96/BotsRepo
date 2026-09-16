import json
from pathlib import Path
from typing import List

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from .models import RawMessage, ScoredMessage


def _dump_json(rows: List[dict], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")


def save_raw_json(messages: List[RawMessage], path: Path) -> None:
    _dump_json([
        {
            "id": m.id,
            "user_id": m.user_id,
            "username": m.username,
            "first_name": m.first_name,
            "last_name": m.last_name,
            "text": m.text,
            "date": m.date.isoformat(),
            "is_bot": m.is_bot,
        }
        for m in messages
    ], path)


def save_filtered_json(messages: List[ScoredMessage], path: Path) -> None:
    _dump_json([
        {
            "user_id": m.user_id,
            "username": m.username,
            "name": m.name,
            "source": m.source,
            "text": m.text,
            "score": m.score,
            "date": m.date.isoformat(),
        }
        for m in messages
    ], path)


def _write_text_cell(ws, row: int, col: int, value, font: Font):
    """Write a string cell, forcing text type so a value like '=...' or '@...'
    is never interpreted by Excel as a formula (CSV/formula-injection guard).
    User-controlled fields (name, username, quote) flow through here."""
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str):
        cell.data_type = "s"
    cell.font = font
    return cell


def save_xlsx(messages: List[ScoredMessage], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Leads"

    headers = ["ID", "Name", "Username", "Группа", "Цитата",
               "Балл готовности (0-100)", "Бэнд", "Обоснование", "Сценарий входа"]
    header_font = Font(name="Arial", bold=True, size=10)
    header_fill = PatternFill("solid", start_color="FFD9E1F2")

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    col_widths = [12, 25, 20, 20, 60, 16, 16, 40, 50]
    for col_idx, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    row_font = Font(name="Arial", size=10)
    for row_idx, msg in enumerate(messages, 2):
        # Numeric columns stay numeric; every user-controlled string goes through
        # _write_text_cell so leading '='/'+'/'-'/'@' can't become a live formula.
        ws.cell(row=row_idx, column=1, value=msg.user_id).font = row_font   # ID
        _write_text_cell(ws, row_idx, 2, msg.name, row_font)                # Name
        _write_text_cell(ws, row_idx, 3, msg.username, row_font)            # Username
        _write_text_cell(ws, row_idx, 4, msg.source, row_font)             # Группа
        _write_text_cell(ws, row_idx, 5, msg.text, row_font)               # Цитата
        # Балл готовности (0-100) · Бэнд · Обоснование · Сценарий входа заполняет Claude —
        # это смысловая оценка по docs/lead-readiness-scoring.md, код её не считает.
        # Предфильтр-score (0–10) сюда НЕ выводится, но остаётся в messages_filtered.json.
        _write_text_cell(ws, row_idx, 6, "", row_font)                     # Балл готовности (0-100)
        _write_text_cell(ws, row_idx, 7, "", row_font)                     # Бэнд
        _write_text_cell(ws, row_idx, 8, "", row_font)                     # Обоснование
        _write_text_cell(ws, row_idx, 9, "", row_font)                     # Сценарий входа

    ws.freeze_panes = "A2"
    wb.save(str(path))


def print_summary(raw_count: int, messages: List[ScoredMessage], filtered_path: Path) -> None:
    print(f"\nСобрано: {raw_count} сообщений → после фильтра: {len(messages)} кандидатов")
    print(f"Сохранено: {filtered_path}\n")

    if not messages:
        print("⚠ Кандидаты не найдены. Попробуйте снизить --min-score.")
        return

    header = (f"{'#':<4} {'ID':<12} {'Name':<18} {'Username':<16} "
              f"{'Группа':<16} {'Предфильтр':<10} Цитата")
    print(header)
    print("-" * 110)
    for idx, msg in enumerate(messages, 1):
        quote = msg.text[:60].replace("\n", " ")
        if len(msg.text) > 60:
            quote += "..."
        print(f"{idx:<4} {msg.user_id:<12} {msg.name[:16]:<18} "
              f"{msg.username[:14]:<16} {msg.source[:14]:<16} {msg.score:<10} {quote}")
