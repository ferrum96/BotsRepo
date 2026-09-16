"""
Генератор примера результата — ТОЛЬКО ВЫМЫШЛЕННЫЕ ДАННЫЕ.
Запускать: python examples/make_sample.py
Создаёт: examples/sample_leads.xlsx
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# Путь к файлу рядом со скриптом
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_PATH = os.path.join(SCRIPT_DIR, "sample_leads.xlsx")

# Полностью вымышленные строки — никаких реальных людей или чатов
ROWS = [
    {
        "#": 1,
        "Имя": "Пример Один",
        "@username": "@example_user1",
        "Чат": "@your_niche_chat",
        "Запрос": "Ищу подрядчика по контенту, какой примерно бюджет?",
        "Балл": 74,
        "Бэнд": "🔥 Горячий",
        "Риск": "🟢 Низкий",
        "Почему": "Активный поиск исполнителя + спрашивает о цене",
        "Что предложить": "Прямой персональный оффер с примерами работ",
    },
    {
        "#": 2,
        "Имя": "Пример Два",
        "@username": "@example_user2",
        "Чат": "@your_niche_chat",
        "Запрос": "Нужен кто-то срочно, запуск через неделю, куда писать?",
        "Балл": 87,
        "Бэнд": "🌋 Очень горячий",
        "Риск": "🟢 Низкий",
        "Почему": "Срочность + конкретный дедлайн + просит контакт",
        "Что предложить": "Связаться немедленно, предложить быстрый старт",
    },
    {
        "#": 3,
        "Имя": "Пример Три",
        "@username": "@example_user3",
        "Чат": "@another_sample_chat",
        "Запрос": "Видел несколько вариантов, сравниваю по срокам и кейсам",
        "Балл": 67,
        "Бэнд": "🔥 Горячий",
        "Риск": "🟡 Проверить вручную",
        "Почему": "Оценивает и сравнивает — стадия выбора исполнителя",
        "Что предложить": "Показать кейсы, дать конкретные сроки и цену",
    },
    {
        "#": 4,
        "Имя": "Пример Четыре",
        "@username": "@example_user4",
        "Чат": "@your_niche_chat",
        "Запрос": "Интересно, как вообще делают такой контент",
        "Балл": 22,
        "Бэнд": "🌤 Тёплый-низкий",
        "Риск": "🟢 Низкий",
        "Почему": "Общий интерес без личной боли или конкретного запроса",
        "Что предложить": "Прогрев контентом, без прямого оффера",
    },
    {
        "#": 5,
        "Имя": "Пример Пять",
        "@username": "@example_user5",
        "Чат": "@another_sample_chat",
        "Запрос": "Ищу монтажёра для коротких видео, бюджет есть",
        "Балл": 58,
        "Бэнд": "☀️ Тёплый",
        "Риск": "🟢 Низкий",
        "Почему": "Активный поиск + упомянул бюджет, но не называет сроки",
        "Что предложить": "Мягкий заход через пользу — показать пример работы",
    },
    {
        "#": 6,
        "Имя": "Пример Шесть",
        "@username": "@example_user6",
        "Чат": "@your_niche_chat",
        "Запрос": "Продаю курс по созданию контента, смотрите в профиле",
        "Балл": 5,
        "Бэнд": "❄️ Холодный/нецелевой",
        "Риск": "🔴 Фрод",
        "Почему": "Сам продаёт / конкурент — штраф −30; рекламный спам",
        "Что предложить": "Не контактировать",
    },
]

HEADERS = ["#", "Имя", "@username", "Чат", "Запрос", "Балл", "Бэнд", "Риск", "Почему", "Что предложить"]

# Цвета бэндов (фон строки)
BAND_COLORS = {
    "🌋 Очень горячий": "FF4500",
    "🔥 Горячий":       "FF8C00",
    "☀️ Тёплый":        "FFD700",
    "🌤 Тёплый-низкий": "ADD8E6",
    "❄️ Холодный/нецелевой": "D3D3D3",
}

HEADER_FILL = PatternFill("solid", fgColor="2F4F4F")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def make_sample():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Примеры лидов"

    # Заголовок
    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Данные
    for row_idx, row_data in enumerate(ROWS, start=2):
        band = row_data["Бэнд"]
        row_color = BAND_COLORS.get(band, "FFFFFF")
        fill = PatternFill("solid", fgColor=row_color)

        for col_idx, header in enumerate(HEADERS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row_data[header])
            cell.fill = fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # Ширина колонок
    col_widths = {
        "#": 4,
        "Имя": 16,
        "@username": 18,
        "Чат": 22,
        "Запрос": 45,
        "Балл": 7,
        "Бэнд": 20,
        "Риск": 20,
        "Почему": 40,
        "Что предложить": 40,
    }
    for col_idx, header in enumerate(HEADERS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = col_widths.get(header, 15)

    ws.row_dimensions[1].height = 30

    # Сноска
    note_row = len(ROWS) + 3
    ws.cell(row=note_row, column=1,
            value="⚠️ Все данные в этом файле полностью вымышлены. Это пример формата вывода инструмента.")
    ws.cell(row=note_row, column=1).font = Font(italic=True, color="888888")
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=len(HEADERS))

    wb.save(OUTPUT_PATH)
    print(f"Файл сохранён: {OUTPUT_PATH}")
    print(f"Строк данных: {len(ROWS)} (+ 1 заголовок)")


if __name__ == "__main__":
    make_sample()
