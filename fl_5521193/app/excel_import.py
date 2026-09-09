"""Excel (.xlsx) import for partner leads."""

from __future__ import annotations

from io import BytesIO
from typing import Any, Optional

from openpyxl import Workbook, load_workbook

from .models import ParsedLead

# Canonical column → accepted header aliases (lower)
COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "first_name": ("first_name", "имя", "name", "firstname"),
    "last_name": ("last_name", "фамилия", "lastname", "surname"),
    "telegram_username": (
        "telegram_username",
        "telegram",
        "username",
        "tg",
        "tg_username",
        "телеграм",
        "telegram username",
    ),
    "telegram_id": ("telegram_id", "tg_id", "telegram id", "id telegram"),
    "phone": ("phone", "телефон", "tel", "mobile"),
    "whatsapp": ("whatsapp", "wa", "ватсап"),
    "email": ("email", "e-mail", "почта", "mail"),
    "tg_channel": ("tg_channel", "channel", "канал", "telegram_channel"),
    "vk": ("vk", "вк", "vkontakte"),
    "site": ("site", "сайт", "website", "url"),
    "instagram": ("instagram", "ig", "инстаграм"),
    "school": ("school", "школа", "проект", "project"),
    "city": ("city", "город", "city/country", "город/страна"),
    "source_found": ("source_found", "источник", "source"),
    "comment": ("comment", "комментарий", "notes"),
    "is_vedic": ("is_vedic", "ведический", "джйотиш", "vedic"),
    "activity_note": ("activity_note", "деятельность", "activity"),
}

TEMPLATE_HEADERS = [
    "first_name",
    "last_name",
    "telegram_username",
    "telegram_id",
    "phone",
    "whatsapp",
    "email",
    "tg_channel",
    "vk",
    "site",
    "instagram",
    "school",
    "city",
    "source_found",
    "comment",
    "is_vedic",
    "activity_note",
]


def _norm_header(h: Any) -> str:
    return str(h or "").strip().lower()


def _map_headers(headers: list[Any]) -> dict[str, int]:
    """field → column index."""
    mapping: dict[str, int] = {}
    normalized = [_norm_header(h) for h in headers]
    for field, aliases in COLUMN_ALIASES.items():
        for idx, h in enumerate(normalized):
            if h in aliases:
                mapping[field] = idx
                break
    return mapping


def _cell(row: tuple[Any, ...], idx: Optional[int]) -> Optional[str]:
    if idx is None or idx >= len(row):
        return None
    val = row[idx]
    if val is None:
        return None
    if isinstance(val, float) and val.is_integer():
        val = int(val)
    s = str(val).strip()
    return s or None


def _boolish(val: Optional[str]) -> bool:
    if val is None:
        return True
    return val.strip().lower() in {"1", "true", "yes", "да", "y", "истина"}


def parse_excel(data: bytes) -> tuple[list[ParsedLead], list[dict]]:
    """Return (leads, row_errors)."""
    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        return [], [{"row": 0, "error": "Пустой файл"}]

    colmap = _map_headers(list(header_row))
    if "first_name" not in colmap:
        return [], [{"row": 1, "error": "Нет колонки first_name / имя"}]

    leads: list[ParsedLead] = []
    errors: list[dict] = []
    for i, row in enumerate(rows, start=2):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        first = _cell(row, colmap.get("first_name"))
        if not first:
            errors.append({"row": i, "error": "Пустое имя"})
            continue
        tg_id_raw = _cell(row, colmap.get("telegram_id"))
        tg_id = None
        if tg_id_raw:
            digits = "".join(ch for ch in tg_id_raw if ch.isdigit())
            tg_id = int(digits) if digits else None

        uname = _cell(row, colmap.get("telegram_username"))
        if uname:
            uname = uname.lstrip("@")

        try:
            leads.append(
                ParsedLead(
                    first_name=first,
                    last_name=_cell(row, colmap.get("last_name")) or "",
                    telegram_username=uname,
                    telegram_id=tg_id,
                    phone=_cell(row, colmap.get("phone")),
                    whatsapp=_cell(row, colmap.get("whatsapp")),
                    email=_cell(row, colmap.get("email")),
                    tg_channel=_cell(row, colmap.get("tg_channel")),
                    vk=_cell(row, colmap.get("vk")),
                    site=_cell(row, colmap.get("site")),
                    instagram=_cell(row, colmap.get("instagram")),
                    school=_cell(row, colmap.get("school")),
                    city=_cell(row, colmap.get("city")),
                    source_found=_cell(row, colmap.get("source_found")) or "Excel",
                    comment=_cell(row, colmap.get("comment")),
                    is_vedic=_boolish(_cell(row, colmap.get("is_vedic"))),
                    activity_note=_cell(row, colmap.get("activity_note")),
                )
            )
        except Exception as exc:  # noqa: BLE001
            errors.append({"row": i, "error": str(exc)})
    wb.close()
    return leads, errors


def build_template_xlsx(sample_leads: Optional[list[ParsedLead]] = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "leads"
    ws.append(TEMPLATE_HEADERS)
    if sample_leads:
        for lead in sample_leads:
            ws.append(
                [
                    lead.first_name,
                    lead.last_name,
                    lead.telegram_username,
                    lead.telegram_id,
                    lead.phone,
                    lead.whatsapp,
                    lead.email,
                    lead.tg_channel,
                    lead.vk,
                    lead.site,
                    lead.instagram,
                    lead.school,
                    lead.city,
                    lead.source_found,
                    lead.comment,
                    "да" if lead.is_vedic else "нет",
                    lead.activity_note,
                ]
            )
    else:
        ws.append(
            [
                "Арина",
                "Белова",
                "astro_partner_01",
                500001,
                "+79032000001",
                "+79032000001",
                "astro_partner_01@mail.example",
                "",
                "",
                "",
                "",
                "Jyotish Lab",
                "Москва",
                "Telegram",
                "пример строки",
                "да",
                "практикующий джйотиш",
            ]
        )
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
